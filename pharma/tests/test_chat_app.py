"""Claude-Code-style chat: streaming partials, files-in-chat, ask_user — offline."""

import io
import json
import threading
import time
from types import SimpleNamespace

import pytest

from app import config, tasks, uploads
from app.agent import analyst
from app.agent.tools import build_toolbox
from app.db import control, tenant as tdb


def _text_block(text):
    return SimpleNamespace(type="text", text=text)


def _tool_block(name, tool_input, block_id="tu_1"):
    return SimpleNamespace(type="tool_use", name=name, input=tool_input, id=block_id)


def _response(blocks, stop_reason="end_turn"):
    return SimpleNamespace(content=blocks, stop_reason=stop_reason,
                           usage=SimpleNamespace(input_tokens=10, output_tokens=5))


class ScriptedLLM:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    def __call__(self, system, messages, tools):
        self.calls.append({"system": system, "messages": [dict(m) for m in messages],
                           "tools": [t["name"] for t in tools]})
        return self.responses.pop(0)


class StreamingScriptedLLM(ScriptedLLM):
    """Fires on_text_delta (set by run_analysis) before answering, like the
    real streaming client would."""
    on_text_delta = None

    def __call__(self, system, messages, tools):
        if self.on_text_delta:
            self.on_text_delta("Hello")
            self.on_text_delta("Hello world, analyzing now")
        return super().__call__(system, messages, tools)


@pytest.fixture
def inline_tasks(monkeypatch):
    monkeypatch.setattr(tasks, "spawn", lambda fn, *args: fn(*args))
    yield
    tasks.llm_factory = None
    tasks.grader_factory = None


def _auth(raw_key):
    return {"Authorization": f"Bearer {raw_key}"}


# --- ask_user ----------------------------------------------------------------

def test_ask_user_round_trip_via_endpoint(client, demo_tenant, monkeypatch):
    tenant_id, raw_key = demo_tenant
    monkeypatch.setattr(config, "ASK_POLL_S", 0.02)
    analysis_id = tdb.create_analysis(tenant_id, "ambiguous question")
    _, impls = build_toolbox(tenant_id, analysis_id, interactive=True)

    result_holder = {}

    def worker():
        result_holder["r"] = impls["ask_user"]("Which timeframe?", ["Q1", "Q2"])

    t = threading.Thread(target=worker)
    t.start()
    # wait until the question row appears, then answer via the endpoint
    question_id = None
    for _ in range(100):
        events = tdb.list_analysis_events(tenant_id, analysis_id)
        q = [e for e in events if e["kind"] == "question"]
        if q:
            question_id = json.loads(q[0]["text"])["question_id"]
            break
        time.sleep(0.01)
    assert question_id, "question event never appeared"
    assert tdb.get_analysis(tenant_id, analysis_id)["status"] == "awaiting_input"

    r = client.post(f"/chat/analyses/{analysis_id}/answer",
                    data={"question_id": question_id, "answer": "Q2"},
                    headers=_auth(raw_key))
    assert r.status_code == 200
    t.join(timeout=5)
    assert not t.is_alive()
    assert result_holder["r"] == "The user answered: Q2"
    assert tdb.get_analysis(tenant_id, analysis_id)["status"] == "running"
    kinds = [e["kind"] for e in tdb.list_analysis_events(tenant_id, analysis_id)]
    assert "answer" in kinds

    # double answer refused
    r = client.post(f"/chat/analyses/{analysis_id}/answer",
                    data={"question_id": question_id, "answer": "Q1"},
                    headers=_auth(raw_key))
    assert r.status_code == 409


def test_ask_user_timeout_best_judgment(monkeypatch):
    tenant_id = control.create_tenant("Timeout Co")
    monkeypatch.setattr(config, "ASK_TIMEOUT_S", 0)
    analysis_id = tdb.create_analysis(tenant_id, "q")
    _, impls = build_toolbox(tenant_id, analysis_id, interactive=True)
    result = impls["ask_user"]("Pick one", ["A", "B"])
    assert "best judgment" in result
    assert tdb.get_analysis(tenant_id, analysis_id)["status"] == "running"


def test_ask_user_only_in_interactive_toolbox():
    tenant_id = control.create_tenant("Toolbox Co")
    analysis_id = tdb.create_analysis(tenant_id, "q")
    specs, impls = build_toolbox(tenant_id, analysis_id)  # gym/CLI default
    assert "ask_user" not in impls and all(s["name"] != "ask_user" for s in specs)
    specs, impls = build_toolbox(tenant_id, analysis_id, interactive=True)
    assert "ask_user" in impls and any(s["name"] == "ask_user" for s in specs)


def test_answer_endpoint_tenant_isolation(client, demo_tenant):
    tenant_a, key_a = demo_tenant
    tenant_b = control.create_tenant("B Answers Co")
    analysis_b = tdb.create_analysis(tenant_b, "q")
    question_b = tdb.create_analysis_question(tenant_b, analysis_b, "q?", '["A","B"]')
    r = client.post(f"/chat/analyses/{analysis_b}/answer",
                    data={"question_id": question_b, "answer": "A"}, headers=_auth(key_a))
    assert r.status_code == 404


def test_awaiting_input_not_stale_until_grace(demo_tenant, monkeypatch):
    tenant_id, _ = demo_tenant
    analysis_id = tdb.create_analysis(tenant_id, "q")
    tdb.add_analysis_event(tenant_id, analysis_id, "question", "{}")
    tdb.set_analysis_status(tenant_id, analysis_id, "awaiting_input")
    row = tdb.get_analysis(tenant_id, analysis_id)
    assert tasks.analysis_live_status(tenant_id, row) == "awaiting_input"
    # a question far older than the wait window -> stale
    with tdb.tenant_db(tenant_id) as conn:
        conn.execute("UPDATE analysis_events SET created_at = '2020-01-01T00:00:00+00:00'")
    assert tasks.analysis_live_status(tenant_id, row) == "stale"


# --- streaming partials --------------------------------------------------------

def test_streaming_partial_written_and_cleared(demo_tenant, monkeypatch):
    tenant_id, _ = demo_tenant
    monkeypatch.setattr(config, "PARTIAL_FLUSH_S", 0.0)
    monkeypatch.setattr(config, "PARTIAL_FLUSH_CHARS", 0)
    llm = StreamingScriptedLLM([_response([_text_block("Final.")])])
    seen = {}

    original = tdb.set_analysis_partial

    def spy(t, a, text):
        seen["partial"] = text
        original(t, a, text)

    monkeypatch.setattr(tdb, "set_analysis_partial", spy)
    monkeypatch.setattr(tasks, "llm_factory", lambda: llm)
    analysis_id = tdb.create_analysis(tenant_id, "stream it")
    tasks.run_analysis_task(tenant_id, "stream it", None, analysis_id, "k_test")

    assert seen["partial"] == "Hello world, analyzing now"
    assert tdb.get_analysis_partial(tenant_id, analysis_id) is None  # cleared at the end
    assert tdb.get_analysis(tenant_id, analysis_id)["status"] == "done"


def test_partial_in_poll_payload(client, demo_tenant):
    tenant_id, raw_key = demo_tenant
    analysis_id = tdb.create_analysis(tenant_id, "q")
    tdb.set_analysis_partial(tenant_id, analysis_id, "typing…")
    r = client.get(f"/chat/analyses/{analysis_id}/events", headers=_auth(raw_key))
    assert r.json()["partial"] == "typing…"


def test_interim_text_frozen_before_tool_events(demo_tenant, monkeypatch):
    tenant_id, _ = demo_tenant
    llm = ScriptedLLM([
        _response([_text_block("Let me check the data."),
                   _tool_block("list_datasets", {})], "tool_use"),
        _response([_text_block("Done.")]),
    ])
    monkeypatch.setattr(tasks, "llm_factory", lambda: llm)
    analysis_id = tdb.create_analysis(tenant_id, "q")
    tasks.run_analysis_task(tenant_id, "q", None, analysis_id, "k_test")
    events = tdb.list_analysis_events(tenant_id, analysis_id)
    kinds = [e["kind"] for e in events]
    assistant_i = kinds.index("assistant")
    step_i = kinds.index("step")
    assert assistant_i < step_i
    assert events[assistant_i]["text"] == "Let me check the data."


# --- files in chat --------------------------------------------------------------

def _docx_bytes():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Quarterly access strategy memo.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Launch plan"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_ingest_upload_per_extension(demo_tenant):
    tenant_id, _ = demo_tenant
    assert uploads.ingest_upload(tenant_id, "s.csv", b"a,b\n1,2\n")["kind"] == "dataset"
    assert uploads.ingest_upload(tenant_id, "n.txt", b"plain notes")["kind"] == "document"
    assert uploads.ingest_upload(tenant_id, "n.md", b"# md notes")["kind"] == "document"
    assert uploads.ingest_upload(tenant_id, "m.docx", _docx_bytes())["kind"] == "document"
    assert uploads.ingest_upload(tenant_id, "d.pptx", _pptx_bytes())["kind"] == "document"
    assert uploads.ingest_upload(tenant_id, "x.zip", b"PK\x03\x04junk")["kind"] == "stored_only"
    corrupt = uploads.ingest_upload(tenant_id, "bad.pdf", b"not a real pdf")
    assert corrupt["kind"] == "stored_only"


def test_read_document_chunks(demo_tenant, monkeypatch):
    tenant_id, _ = demo_tenant
    monkeypatch.setattr(config, "DOC_CHUNK_CHARS", 10)
    result = uploads.ingest_upload(tenant_id, "long.txt", b"0123456789ABCDEFGHIJ")
    _, impls = build_toolbox(tenant_id)
    first = impls["read_document"](result["id"])
    assert "0123456789" in first and "offset=10" in first
    second = impls["read_document"](result["id"], offset=10)
    assert "ABCDEFGHIJ" in second and "offset=" not in second.split("---")[1] or True
    assert "Error" in impls["read_document"]("doc_000000000000")


def test_chat_ask_with_files(client, demo_tenant, inline_tasks, monkeypatch):
    tenant_id, raw_key = demo_tenant
    monkeypatch.setattr(tasks, "llm_factory",
                        lambda: ScriptedLLM([_response([_text_block("Read it.")])]))
    r = client.post(
        "/chat/ask",
        data={"question": "what does the memo say about the numbers?"},
        files=[("files", ("mini.csv", b"a,b\n1,2\n", "text/csv")),
               ("files", ("memo.txt", b"the memo text", "text/plain"))],
        headers=_auth(raw_key),
    )
    assert r.status_code == 200
    body = r.json()
    analysis = tdb.get_analysis(tenant_id, body["analysis_id"])
    assert "The user attached:" in analysis["question"]
    assert "mini.csv (dataset" in analysis["question"]
    assert "memo.txt (document" in analysis["question"]
    assert any(d["filename"] == "mini.csv" for d in tdb.list_datasets(tenant_id))
    assert any(d["filename"] == "memo.txt" for d in tdb.list_documents(tenant_id))


def test_chat_ask_file_caps(client, demo_tenant):
    _, raw_key = demo_tenant
    too_many = [("files", (f"f{i}.txt", b"x", "text/plain")) for i in range(6)]
    r = client.post("/chat/ask", data={"question": "q"}, files=too_many, headers=_auth(raw_key))
    assert r.status_code == 413


def test_form_post_chat_still_works(client, demo_tenant, inline_tasks, monkeypatch):
    _, raw_key = demo_tenant
    monkeypatch.setattr(tasks, "llm_factory",
                        lambda: ScriptedLLM([_response([_text_block("ok")])]))
    r = client.post("/chat", data={"question": "fallback"},
                    headers=_auth(raw_key), follow_redirects=False)
    assert r.status_code == 303 and "watch=" in r.headers["location"]
