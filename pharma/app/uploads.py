"""One ingest path for every file a client hands us.

csv/xlsx become datasets (the agent computes over them in the sandbox);
pdf/txt/md/docx/pptx become documents (text extracted, readable via the
read_document tool); anything else is stored with an honest "can't read this
yet". Extraction failures degrade to stored-only — never a crash.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from pathlib import Path

from app.db import tenant as tdb

_SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._-]+")

DATASET_EXTS = (".csv", ".xlsx")
DOCUMENT_EXTS = (".pdf", ".txt", ".md", ".docx", ".pptx")


def sanitize_filename(filename: str | None, default: str = "upload") -> str:
    return _SAFE_NAME_RE.sub("_", filename or default).lstrip(".") or default


def extract_document_text(filename: str, data: bytes) -> str | None:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip() or None
        if suffix in (".txt", ".md"):
            return data.decode(errors="replace").strip() or None
        if suffix == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs).strip() or None
        if suffix == ".pptx":
            from pptx import Presentation

            prs = Presentation(io.BytesIO(data))
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                texts = [shape.text_frame.text for shape in slide.shapes
                         if shape.has_text_frame and shape.text_frame.text.strip()]
                if texts:
                    slides.append(f"[Slide {i}]\n" + "\n".join(texts))
            return "\n\n".join(slides).strip() or None
    except Exception:
        return None
    return None


def _ingest_dataset(tenant_id: str, filename: str, data: bytes) -> dict:
    rows = None
    schema_json = None
    try:
        import pandas as pd

        df = pd.read_csv(io.BytesIO(data)) if filename.lower().endswith(".csv") \
            else pd.read_excel(io.BytesIO(data))
        rows = len(df)
        schema_json = json.dumps({c: str(t) for c, t in df.dtypes.items()})
    except Exception:
        pass  # store anyway; the agent can inspect it
    dest = tdb.resolve_tenant_file(tenant_id, f"uploads/{filename}")
    dest.write_bytes(data)
    dataset_id = tdb.add_dataset(
        tenant_id, filename, f"uploads/{filename}",
        hashlib.sha256(data).hexdigest(), rows, schema_json,
    )
    note = f"{filename} (dataset {dataset_id}" + (f", {rows} rows)" if rows is not None else ")")
    return {"kind": "dataset", "id": dataset_id, "filename": filename, "rows": rows, "note": note}


def ingest_upload(tenant_id: str, filename: str | None, data: bytes) -> dict:
    filename = sanitize_filename(filename)
    suffix = Path(filename).suffix.lower()

    if suffix in DATASET_EXTS:
        return _ingest_dataset(tenant_id, filename, data)

    dest = tdb.resolve_tenant_file(tenant_id, f"uploads/{filename}")
    dest.write_bytes(data)
    text = extract_document_text(filename, data) if suffix in DOCUMENT_EXTS else None
    document_id = tdb.add_document(tenant_id, filename, f"uploads/{filename}", text)
    if text:
        note = f"{filename} (document {document_id}, {len(text)} chars extracted)"
        return {"kind": "document", "id": document_id, "filename": filename, "note": note}
    note = f"{filename} (stored as {document_id}, but I can't read this file type yet)"
    return {"kind": "stored_only", "id": document_id, "filename": filename, "note": note}
