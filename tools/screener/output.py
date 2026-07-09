"""tools/screener/output.py — scored XLSX + summary.md (stage 7).

Sheet ``leads`` carries every processed row (all original input columns,
Phone/DNC untouched, then the enrichment columns); sheet ``needs_manual``
holds the rows the pipeline refused to judge, with the reason. Whole-row
fills encode the verdict so the sheet reads at a glance:
green=SEND CONTRACT, yellow=NEGOTIATE, orange=ASSEMBLAGE_LEAD/RENEGOTIATE,
red=PASS.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

NEW_COLUMNS = (
    "hcad_account", "width_ft", "depth_ft", "aspect_ratio", "shape_flag",
    "frontage", "frontage_street", "flood_zone", "flood_flag", "street_ppsf",
    "retail_estimate", "buyer_ceiling", "mao", "verdict", "open_at",
    "walk_at", "adjacent_owners", "comp_evidence", "needs_manual_reason",
)

VERDICT_FILLS = {
    "SEND CONTRACT": "C6EFCE",   # green
    "NEGOTIATE": "FFEB9C",       # yellow
    "ASSEMBLAGE_LEAD": "F8CBAD", # orange
    "RENEGOTIATE": "F8CBAD",     # orange
    "PASS": "FFC7CE",            # red
}

_MONEY_COLUMNS = {"retail_estimate", "buyer_ceiling", "mao", "open_at", "walk_at"}


def _write_sheet(ws, rows: list[dict], headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append([row.get(h) for h in headers])
        fill_hex = VERDICT_FILLS.get(row.get("verdict") or "")
        if fill_hex:
            fill = PatternFill(start_color=fill_hex, end_color=fill_hex,
                               fill_type="solid")
            for cell in ws[ws.max_row]:
                cell.fill = fill
        for idx, h in enumerate(headers, start=1):
            if h in _MONEY_COLUMNS:
                ws.cell(row=ws.max_row, column=idx).number_format = "#,##0"

    if rows:
        ws.auto_filter.ref = ws.dimensions


def write_xlsx(rows: list[dict], input_headers: list[str], path: str | Path) -> None:
    """rows = all rows (processed + needs_manual mixed); split by reason."""

    headers = list(input_headers) + [c for c in NEW_COLUMNS if c not in input_headers]
    processed = [r for r in rows if not r.get("needs_manual_reason")]
    manual = [r for r in rows if r.get("needs_manual_reason")]

    wb = Workbook()
    _write_sheet(wb.active, processed, headers)
    wb.active.title = "leads"
    _write_sheet(wb.create_sheet("needs_manual"), manual, headers)
    wb.save(path)


def write_summary(rows: list[dict], path: str | Path) -> None:
    processed = [r for r in rows if not r.get("needs_manual_reason")]
    manual = [r for r in rows if r.get("needs_manual_reason")]

    counts: dict[str, int] = {}
    for r in processed:
        v = r.get("verdict") or "(unscored)"
        counts[v] = counts.get(v, 0) + 1

    any_asking = any(r.get("_asking") is not None for r in processed)

    def rank_key(r: dict) -> float:
        retail = r.get("retail_estimate") or 0
        if any_asking and r.get("_asking") is not None:
            return retail - r["_asking"]
        return retail

    callable_rows = [r for r in processed if r.get("verdict") not in ("PASS", "")]
    top = sorted(callable_rows, key=rank_key, reverse=True)[:10]

    lines = ["# Screening summary", "", "## Verdicts", ""]
    for verdict in ("SEND CONTRACT", "NEGOTIATE", "RENEGOTIATE",
                    "ASSEMBLAGE_LEAD", "PASS"):
        if verdict in counts:
            lines.append(f"- **{verdict}**: {counts.pop(verdict)}")
    for verdict, n in sorted(counts.items()):
        lines.append(f"- **{verdict}**: {n}")
    lines.append(f"- needs_manual: {len(manual)}")

    ranked_by = "retail − asking" if any_asking else "retail estimate"
    lines += ["", f"## Top 10 leads (by {ranked_by})", ""]
    if top:
        lines.append("| Address | Zip | Verdict | Retail | MAO | Asking |")
        lines.append("|---|---|---|---|---|---|")
        for r in top:
            def money(v):
                return f"${v:,.0f}" if isinstance(v, (int, float)) else "—"
            lines.append(
                f"| {r.get('Address') or '?'} | {r.get('Zip') or ''} "
                f"| {r.get('verdict')} | {money(r.get('retail_estimate'))} "
                f"| {money(r.get('mao'))} | {money(r.get('_asking'))} |"
            )
    else:
        lines.append("(no callable leads)")

    Path(path).write_text("\n".join(lines) + "\n", encoding="utf-8")
