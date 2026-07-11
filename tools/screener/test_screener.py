"""Offline tests for the Dirt Screener — every seam stubbed, no network."""

from __future__ import annotations

import json
import math

from .cache import Cache
from .comps import (
    BraveAuthError,
    BraveClient,
    build_queries,
    extract_comps,
    snippet_quality,
    street_name,
)
from .config import DEFAULT_CONFIG, config_hash, load_config
from .fema import classify_zone, flood_zone
from .frontage import compute_frontage, fetch_roads
from .geometry import (
    from_local_feet,
    lot_mismatch,
    parcel_metrics,
    shape_flag,
    to_local_feet,
)
from .harris import fetch_adjacent, fetch_parcel, normalize_apn
from .arcgis import query_layer
from .output import NEW_COLUMNS, VERDICT_FILLS, write_summary, write_xlsx
from .scoring import BUILDER_RE, assemblage_candidates, score_row

LON0, LAT0 = -95.28, 29.851  # Houston 77028


# ---------------------------------------------------------------------------
# fixtures
# ---------------------------------------------------------------------------


def _rect_rings(width_ft: float, depth_ft: float, theta_deg: float = 0.0):
    """A width x depth rectangle rotated theta, as WGS84 rings at Houston."""

    t = math.radians(theta_deg)
    corners = [(0, 0), (depth_ft, 0), (depth_ft, width_ft), (0, width_ft), (0, 0)]
    rotated = [
        (x * math.cos(t) - y * math.sin(t), x * math.sin(t) + y * math.cos(t))
        for x, y in corners
    ]
    return from_local_feet([rotated], LON0, LAT0)


SLIVER_RINGS = _rect_rings(33, 660, theta_deg=30)   # the brief's must-catch strip
NORMAL_RINGS = _rect_rings(60, 120)

HCAD_FEATURE = {
    "features": [{
        "attributes": {"HCAD_NUM": "0440240000280", "owner_name_1": "JANE DOE",
                       "Acreage": "0.5"},
        "geometry": {"rings": [[list(pt) for pt in SLIVER_RINGS[0]]]},
    }]
}
HCAD_ADJ = {
    "features": [
        {"attributes": {"HCAD_NUM": "0440240000281",
                        "owner_name_1": "LONE STAR HOMES LLC"}},
        {"attributes": {"HCAD_NUM": "0440240000282",
                        "owner_name_1": "LONE STAR HOMES LLC"}},
        {"attributes": {"HCAD_NUM": "0440240000283",
                        "owner_name_1": "SMITH JOHN"}},
        {"attributes": {"HCAD_NUM": "0440240000280",  # self — must be dropped
                        "owner_name_1": "JANE DOE"}},
    ]
}
FEMA_AE = {"features": [{"attributes": {"FLD_ZONE": "AE", "ZONE_SUBTY": "",
                                        "SFHA_TF": "T"}}]}
FEMA_X = {"features": [{"attributes": {"FLD_ZONE": "X", "ZONE_SUBTY": "",
                                       "SFHA_TF": "F"}}]}

# a road 2 ft north of the normal lot's top edge (inside the 5 ft buffer)
_road_ft = [(-50.0, 62.0), (170.0, 62.0)]
TXDOT_ROADS = {
    "features": [{
        "attributes": {"ST_NM": "Richland Dr", "RTE_NM": ""},
        "geometry": {"paths": [
            [list(pt) for pt in from_local_feet([_road_ft], LON0, LAT0)[0]]
        ]},
    }]
}
OVERPASS_ROAD = {
    "elements": [{
        "type": "way",
        "tags": {"name": "Richland Dr", "highway": "residential"},
        "geometry": [
            {"lon": lon, "lat": lat}
            for lon, lat in from_local_feet([_road_ft], LON0, LAT0)[0]
        ],
    }]
}

BRAVE_GOOD = {
    "web": {"results": [
        {"title": "Lot sold: 4210 Richland Dr", "url": "https://landsite.com/1",
         "description": "0.2 acre lot sold for $45,000 — 8,700 sqft infill"},
        {"title": "77028 land values", "url": "https://comps.com/2",
         "description": "vacant lots trade at $5-8 per sq ft, e.g. $52,000 for 8,000 sqft"},
        {"title": "cute pins", "url": "https://pinterest.com/x",
         "description": "$45,000 lot ideas sqft"},
        {"title": "no numbers here", "url": "https://blog.com/3",
         "description": "land is valuable"},
    ]}
}


class StubLLM:
    def __init__(self, payload: str):
        self.payload = payload
        self.calls = 0

    def invoke(self, prompt: str):
        self.calls += 1
        self.prompt = prompt
        return type("R", (), {"content": self.payload})()


def _arcgis_stub(fixtures: dict):
    """Dispatch stubbed JSON by URL substring; records calls."""

    calls = []

    def stub(method, url, payload, key):
        calls.append(url)
        for marker, fixture in fixtures.items():
            if marker in url:
                return 200, json.dumps(fixture)
        return 200, json.dumps({"features": []})

    return stub, calls


# ---------------------------------------------------------------------------
# geometry
# ---------------------------------------------------------------------------


def test_projection_roundtrip_feet():
    rings = from_local_feet([[(0.0, 0.0), (0.0, 364.6)]], LON0, LAT0)
    (x0, y0), (x1, y1) = to_local_feet(rings, LON0, LAT0)[0]
    assert abs(y1 - 364.6) < 0.01 and abs(x1) < 0.01


def test_sliver_parcel_flagged():
    m = parcel_metrics(SLIVER_RINGS)
    assert abs(m["width_ft"] - 33) < 1 and abs(m["depth_ft"] - 660) < 2
    assert m["aspect_ratio"] > 4
    assert shape_flag(m["width_ft"], m["aspect_ratio"]) == "SLIVER"
    assert abs(m["area_sqft"] - 33 * 660) < 50


def test_normal_lot_passes():
    m = parcel_metrics(NORMAL_RINGS)
    assert abs(m["width_ft"] - 60) < 1 and abs(m["depth_ft"] - 120) < 1
    assert shape_flag(m["width_ft"], m["aspect_ratio"]) is None


def test_wide_but_elongated_flagged():
    m = parcel_metrics(_rect_rings(80, 400))
    assert m["aspect_ratio"] == 5.0
    assert shape_flag(m["width_ft"], m["aspect_ratio"]) == "SLIVER"


def test_degenerate_geometry_none():
    assert parcel_metrics([]) is None
    assert parcel_metrics([[(0, 0), (0, 0), (0, 0), (0, 0)]]) is None


def test_lot_mismatch():
    assert lot_mismatch(10_000, 12_000)          # 20% off
    assert not lot_mismatch(10_000, 10_500)      # 5% off
    assert not lot_mismatch(None, 12_000)


# ---------------------------------------------------------------------------
# arcgis / harris
# ---------------------------------------------------------------------------


def test_normalize_apn():
    assert normalize_apn("044-024-000-0280") == "0440240000280"
    assert normalize_apn("0440240000280") == "0440240000280"
    assert normalize_apn("123") is None
    assert normalize_apn(None) is None
    assert normalize_apn("") is None


def test_query_layer_paginates():
    page1 = {"features": [{"attributes": {"i": 1}}], "exceededTransferLimit": True}
    page2 = {"features": [{"attributes": {"i": 2}}]}
    calls = []

    def stub(method, url, payload, key):
        calls.append(url)
        return 200, json.dumps(page2 if "resultOffset" in url else page1)

    data = query_layer("https://gis/layer/0", {"where": "1=1"},
                       http_request=stub, sleep=lambda s: None)
    assert [f["attributes"]["i"] for f in data["features"]] == [1, 2]
    assert len(calls) == 2


def test_query_layer_retries_then_fails():
    attempts = []

    def stub(method, url, payload, key):
        attempts.append(1)
        return 503, "unavailable"

    data = query_layer("https://gis/layer/0", {}, http_request=stub,
                       sleep=lambda s: None, retry_delays=(0.1, 0.1))
    assert "error" in data and len(attempts) == 3


def test_fetch_parcel_and_adjacent():
    stub, calls = _arcgis_stub({"HCAD_NUM%3D%27044": HCAD_FEATURE,
                                "esriGeometryPolygon": HCAD_ADJ})
    parcel = fetch_parcel("0440240000280", http_request=stub, sleep=lambda s: None)
    assert parcel["owner"] == "JANE DOE" and len(parcel["rings"][0]) == 5

    adj = fetch_adjacent("0440240000280", parcel["rings"],
                         http_request=stub, sleep=lambda s: None)
    accounts = {a["account"] for a in adj}
    assert "0440240000280" not in accounts          # self excluded
    assert len(adj) == 3


def test_fetch_parcel_missing_and_error():
    missing = fetch_parcel("9999999999999",
                           http_request=lambda *a: (200, '{"features": []}'),
                           sleep=lambda s: None)
    assert missing.get("missing")
    err = fetch_parcel("0440240000280",
                       http_request=lambda *a: (500, "boom"), sleep=lambda s: None)
    assert "error" in err


# ---------------------------------------------------------------------------
# putnam (FL) adapter
# ---------------------------------------------------------------------------


def test_putnam_normalize_apn():
    from .putnam import _dashed, normalize_apn

    assert normalize_apn("11-10-23-9303-0020-0230") == "111023930300200230"
    assert normalize_apn("111023930300200230") == "111023930300200230"
    assert normalize_apn("044-024-000-0280") is None      # 13 digits = not FL
    assert normalize_apn(None) is None
    assert _dashed("111023930300200230") == "11-10-23-9303-0020-0230"


def test_putnam_fetch_parcel_retries_dashed_format():
    from .putnam import fetch_parcel

    fl_feature = {
        "features": [{
            "attributes": {"PARCELNO": "11-10-23-9303-0020-0230",
                           "OWN_NAME": "SMITH JANE", "CO_NO": 64},
            "geometry": {"rings": [[list(pt) for pt in NORMAL_RINGS[0]]]},
        }]
    }
    wheres = []

    def stub(method, url, payload, key):
        wheres.append(url)
        # stripped form finds nothing; the dashed retry hits
        if "11-10-23-9303-0020-0230" in url:
            return 200, json.dumps(fl_feature)
        return 200, '{"features": []}'

    parcel = fetch_parcel("111023930300200230", http_request=stub,
                          sleep=lambda s: None)
    assert parcel["owner"] == "SMITH JANE"
    assert len(wheres) == 2                              # stripped, then dashed
    assert all("CO_NO" in u for u in wheres)             # county scoping always


def test_putnam_fetch_adjacent_excludes_self_in_both_formats():
    from .putnam import fetch_adjacent

    adj = {
        "features": [
            {"attributes": {"PARCELNO": "11-10-23-9303-0020-0230",  # self, dashed
                            "OWN_NAME": "SMITH JANE"}},
            {"attributes": {"PARCELNO": "11-10-23-9303-0020-0231",
                            "OWN_NAME": "RIVER DEVELOPMENT LLC"}},
        ]
    }
    out = fetch_adjacent("111023930300200230", NORMAL_RINGS,
                         http_request=lambda *a: (200, json.dumps(adj)),
                         sleep=lambda s: None)
    assert [a["owner"] for a in out] == ["RIVER DEVELOPMENT LLC"]


def test_roads_config_seam_picks_state_layer():
    from . import harris, putnam

    tx_url, tx_fields = harris.roads_config(DEFAULT_CONFIG)
    fl_url, fl_fields = putnam.roads_config(DEFAULT_CONFIG)
    assert "TxDOT" in tx_url and "fdot" in fl_url
    assert "ST_NAME" in fl_fields


# ---------------------------------------------------------------------------
# frontage
# ---------------------------------------------------------------------------


def test_frontage_present_with_street_name():
    roads = [{"name": "Richland Dr",
              "coords": from_local_feet([_road_ft], LON0, LAT0)[0]}]
    result = compute_frontage(NORMAL_RINGS, roads)
    assert result["frontage_street"] == "Richland Dr"
    assert isinstance(result["frontage"], int) and result["frontage"] > 50


def test_frontage_none_when_far():
    far = [{"name": "Far Rd", "coords":
            from_local_feet([[(-50.0, 500.0), (170.0, 500.0)]], LON0, LAT0)[0]}]
    result = compute_frontage(NORMAL_RINGS, far)
    assert result == {"frontage": "NONE", "frontage_street": None}


def test_fetch_roads_backoff_then_fallback_and_failure_is_none():
    seq = iter([(429, "slow down"), (429, ""), (429, ""), (429, ""),
                (200, json.dumps(OVERPASS_ROAD))])
    hosts = []

    def stub(method, url, payload, key):
        hosts.append(url)
        return next(seq)

    bbox = (LON0 - 0.001, LAT0 - 0.001, LON0 + 0.001, LAT0 + 0.001)
    roads = fetch_roads(bbox, http_request=stub, sleep=lambda s: None)
    assert roads and roads[0]["name"] == "Richland Dr"
    assert len(set(hosts)) == 2                     # fell through to fallback

    dead = fetch_roads(bbox, http_request=lambda *a: (429, ""),
                       sleep=lambda s: None)
    assert dead is None                              # never an empty-list kill


def test_fetch_roads_arcgis_parses_paths_and_names():
    from .frontage import fetch_roads_arcgis

    urls = []

    def stub(method, url, payload, key):
        urls.append(url)
        return 200, json.dumps(TXDOT_ROADS)

    bbox = (LON0 - 0.001, LAT0 - 0.001, LON0 + 0.001, LAT0 + 0.001)
    roads = fetch_roads_arcgis(bbox, http_request=stub, sleep=lambda s: None)
    assert roads and roads[0]["name"] == "Richland Dr"
    assert len(roads[0]["coords"]) == 2
    assert "TxDOT_Roadways" in urls[0]

    dead = fetch_roads_arcgis(bbox, http_request=lambda *a: (503, "down"),
                              sleep=lambda s: None)
    assert dead is None                              # failure, not "no roads"

    empty = fetch_roads_arcgis(bbox,
                               http_request=lambda *a: (200, '{"features":[]}'),
                               sleep=lambda s: None)
    assert empty == []                               # verified no roads


def test_overpass_remark_timeout_is_failure_not_no_roads():
    # HTTP 200 + "remark" is Overpass saying it gave up — NOT "no roads here";
    # reading it as empty would false-kill the lead as landlocked.
    overloaded = json.dumps({
        "remark": "runtime error: Query timed out in \"query\" at line 1",
        "elements": [],
    })
    bbox = (LON0 - 0.001, LAT0 - 0.001, LON0 + 0.001, LAT0 + 0.001)
    roads = fetch_roads(bbox, http_request=lambda *a: (200, overloaded),
                        sleep=lambda s: None)
    assert roads is None


# ---------------------------------------------------------------------------
# fema
# ---------------------------------------------------------------------------


def test_classify_zone_matrix():
    assert classify_zone("AE") == "FLOODPLAIN"
    assert classify_zone("A") == "FLOODPLAIN"
    assert classify_zone("VE") == "FLOODPLAIN"
    assert classify_zone("AO") == "FLOODPLAIN"
    assert classify_zone("X") is None
    assert classify_zone("X", "F") is None
    assert classify_zone("X", "T") == "FLOODPLAIN"   # SFHA trumps the letter
    assert classify_zone(None) is None


def test_flood_zone_query():
    result = flood_zone(29.828, -95.286,
                        http_request=lambda *a: (200, json.dumps(FEMA_AE)),
                        sleep=lambda s: None)
    assert result["flood_zone"] == "AE" and result["flood_flag"] == "FLOODPLAIN"
    assert "hazards.fema.gov" in result["source"]


def test_flood_zone_uses_fallback_host():
    urls = []

    def stub(method, url, payload, key):
        urls.append(url)
        if "/gis/nfhl/" in url:  # the fallback path
            return 200, json.dumps(FEMA_AE)
        return 0, "network error: SSL EOF"

    result = flood_zone(29.828, -95.286, http_request=stub, sleep=lambda s: None)
    assert result["flood_flag"] == "FLOODPLAIN"
    assert any("/gis/nfhl/" in u for u in urls)


def test_flood_zone_falls_through_to_agol_mirror():
    # both FEMA hosts down (WAF blocks the VPS) -> the Esri Living Atlas
    # mirror answers; its schema has no SFHA_TF, prefix logic still flags
    agol = {"features": [{"attributes": {"FLD_ZONE": "AE",
                                         "esri_symbology": "1% Annual Chance"}}]}
    urls = []

    def stub(method, url, payload, key):
        urls.append(url)
        if "services.arcgis.com" in url:
            return 200, json.dumps(agol)
        return 0, "network error: SSL EOF"

    result = flood_zone(29.828, -95.286, http_request=stub, sleep=lambda s: None)
    assert result["flood_flag"] == "FLOODPLAIN"
    assert result["flood_zone"] == "AE"
    assert any("services.arcgis.com" in u for u in urls)
    empty = flood_zone(30.0, -95.0,
                       http_request=lambda *a: (200, '{"features": []}'),
                       sleep=lambda s: None)
    assert empty["flood_zone"] == "UNKNOWN" and empty["flood_flag"] is None


def test_sfha_count_in_envelope():
    from .fema import sfha_count_in_envelope

    urls = []

    def stub(method, url, payload, key):
        urls.append(url)
        return 200, '{"count": 42}'

    result = sfha_count_in_envelope((-95.31, 29.80, -95.26, 29.85),
                                    http_request=stub, sleep=lambda s: None)
    assert result["count"] == 42
    assert "returnCountOnly" in urls[0] and "esriGeometryEnvelope" in urls[0]

    err = sfha_count_in_envelope((-95.31, 29.80, -95.26, 29.85),
                                 http_request=lambda *a: (0, "network error: x"),
                                 sleep=lambda s: None)
    assert "error" in err


# ---------------------------------------------------------------------------
# comps
# ---------------------------------------------------------------------------


def test_street_name_and_queries():
    assert street_name("4210 Richland Dr") == "Richland Dr"
    assert street_name("123B Oak St") == "Oak St"
    q = build_queries({"Address": "4210 Richland Dr", "Zip": "77028",
                       "City": "Houston"})
    assert q[0] == '"Richland Dr" 77028 lot sold'
    assert "price per square foot" in q[1]


def test_snippet_quality_gate():
    snips = [
        {"title": "sold $45,000", "description": "8,700 sqft lot", "url": "https://a.com"},
        {"title": "$5,000,000 mansion", "description": "no land info", "url": "https://b.com"},
        {"title": "sold $45,000 sqft", "description": "", "url": "https://pinterest.com/x"},
    ]
    kept = snippet_quality(snips)
    assert len(kept) == 1 and kept[0]["url"] == "https://a.com"


def test_extract_comps_validates_arithmetic():
    good = json.dumps({
        "comps": [
            {"price": 45000, "sqft": 8700, "ppsf": 5.17, "regime": "small",
             "same_street": True, "source": "landsite"},
            {"price": 52000, "sqft": 8000, "ppsf": 6.5, "regime": "small",
             "same_street": False, "source": "comps.com"},
            {"price": 100000, "sqft": 10, "ppsf": 5.0, "regime": "small",
             "same_street": False, "source": "liar"},        # ppsf 10000 vs 5 -> drop
        ],
        "evidence": "two 77028 infill sales",
    })
    llm = StubLLM(good)
    snips = [{"title": f"s{i}", "description": "$45,000 8,700 sqft",
              "url": f"https://x.com/{i}"} for i in range(3)]
    row = {"Address": "4210 Richland Dr", "Zip": "77028", "_lot_sqft": 8000.0}
    out = extract_comps(snips, row, llm=llm)
    assert out["street_ppsf_small"] == 5.84          # median of 5.17, 6.5
    assert out["street_ppsf_acreage"] is None
    assert "UNTRUSTED" in llm.prompt                 # injection guard present


def test_extract_comps_fails_closed():
    llm = StubLLM("I could not find any JSON worth returning")
    snips = [{"title": "a", "description": "$1 sqft", "url": "https://1"},
             {"title": "b", "description": "$2 sqft", "url": "https://2"}]
    out = extract_comps(snips, {"_lot_sqft": 8000.0}, llm=llm)
    assert out["street_ppsf_small"] is None and out["street_ppsf_acreage"] is None

    gate = StubLLM("{}")
    out2 = extract_comps([{"title": "only one", "description": "$1 sqft",
                           "url": "https://1"}], {}, llm=gate)
    assert gate.calls == 0                           # <2 snippets: LLM never called
    assert out2["comp_evidence"] == "insufficient comp snippets"


def test_brave_client_auth_fail_closed(tmp_path):
    client = BraveClient("bad-key", http_request=lambda *a: (401, "denied"),
                         sleep=lambda s: None)
    assert client.search("q1") == []
    assert client.search("q2") == []
    try:
        client.search("q3")
        assert False, "expected BraveAuthError"
    except BraveAuthError:
        pass


def test_brave_client_parses_and_caches(tmp_path):
    cache = Cache(tmp_path / "c.db")
    calls = []

    def stub(method, url, payload, key):
        calls.append(url)
        return 200, json.dumps(BRAVE_GOOD)

    client = BraveClient("k", http_request=stub, sleep=lambda s: None, cache=cache)
    first = client.search("77028 land")
    second = client.search("77028 land")
    assert len(first) == 4 and first == second
    assert len(calls) == 1                           # second hit came from cache


# ---------------------------------------------------------------------------
# scoring
# ---------------------------------------------------------------------------


def test_builder_regex():
    for name in ("LONE STAR HOMES LLC", "ABC Development Inc", "XYZ Builders",
                 "ACME PROPERTIES LTD", "Delta Land Co"):
        assert BUILDER_RE.search(name), name
    for name in ("SMITH JOHN", "MARIA GARCIA", "DOE JANE ESTATE OF"):
        assert not BUILDER_RE.search(name), name


def test_assemblage_candidates():
    adj = [{"owner": "LONE STAR HOMES LLC", "account": "1"},
           {"owner": "SMITH JOHN", "account": "2"},
           {"owner": "SMITH JOHN", "account": "3"}]   # 2 parcels, same person
    hits = assemblage_candidates(adj)
    assert "LONE STAR HOMES LLC" in hits and "SMITH JOHN" in hits


def _scored(**overrides):
    row = {
        "needs_manual_reason": None, "shape_flag": None, "frontage": 60,
        "frontage_street": "Richland Dr", "flood_flag": None, "flood_zone": "X",
        "retail_estimate": 80_000, "_asking": None, "_adjacent": [],
        "comp_evidence": "evidence",
    }
    row.update(overrides)
    return row, score_row(row)


def test_verdict_matrix():
    # NEGOTIATE: ceiling 60k, mao 50k, open 42.5k, walk 55k
    row, out = _scored()
    assert out["verdict"] == "NEGOTIATE"
    assert out["buyer_ceiling"] == 60_000 and out["mao"] == 50_000
    assert out["open_at"] == 42_500 and out["walk_at"] == 55_000

    # SEND CONTRACT when mao >= asking
    _, out = _scored(_asking=45_000)
    assert out["verdict"] == "SEND CONTRACT"

    # asking above mao stays NEGOTIATE
    _, out = _scored(_asking=55_000)
    assert out["verdict"] == "NEGOTIATE"

    # flood outranks SEND CONTRACT and caps valuation 30%
    row, out = _scored(_asking=30_000, flood_flag="FLOODPLAIN", flood_zone="AE")
    assert out["verdict"] == "RENEGOTIATE"
    assert out["buyer_ceiling"] == 42_000            # 80k*0.7*0.75
    assert "30%" in row["comp_evidence"]

    # hard kills
    _, out = _scored(shape_flag="SLIVER")
    assert out["verdict"] == "PASS"
    _, out = _scored(frontage="NONE")
    assert out["verdict"] == "PASS"

    # sliver + builder neighbor -> ASSEMBLAGE_LEAD
    _, out = _scored(shape_flag="SLIVER",
                     _adjacent=[{"owner": "LONE STAR HOMES LLC", "account": "1"}])
    assert out["verdict"] == "ASSEMBLAGE_LEAD"

    # no comps -> needs_manual, blank verdict
    row, out = _scored(retail_estimate=None)
    assert out["verdict"] == "" and row["needs_manual_reason"] == "no usable comps"

    # needs_manual short-circuits
    _, out = _scored(needs_manual_reason="missing APN")
    assert out["verdict"] == ""


# ---------------------------------------------------------------------------
# cache & config
# ---------------------------------------------------------------------------


def test_cache_roundtrip_and_config_hash(tmp_path):
    cache = Cache(tmp_path / "c.db")
    h = config_hash()
    assert cache.get_stage("acct", "geometry", h) is None
    cache.put_stage("acct", "geometry", h, {"width_ft": 33.0})
    assert cache.get_stage("acct", "geometry", h) == {"width_ft": 33.0}
    # a genome change invalidates
    h2 = config_hash(DEFAULT_CONFIG.mutate(min_width_ft=60.0))
    assert h2 != h and cache.get_stage("acct", "geometry", h2) is None
    # non-200s are never cached
    cache.put_http("k", 503, "boom")
    assert cache.get_http("k") is None


def test_load_config_yaml_overlay(tmp_path):
    p = tmp_path / "config.yaml"
    p.write_text("target_fee_usd: 12000\nretry_delays_s: [1.0, 2.0]\n")
    cfg = load_config(str(p))
    assert cfg.target_fee_usd == 12_000 and cfg.retry_delays_s == (1.0, 2.0)
    p.write_text("not_a_knob: 1\n")
    try:
        load_config(str(p))
        assert False, "unknown key must raise"
    except ValueError as exc:
        assert "not_a_knob" in str(exc)


# ---------------------------------------------------------------------------
# output
# ---------------------------------------------------------------------------


def _sample_rows():
    return [
        {"Address": "4210 Richland Dr", "Zip": "77028", "APN": "044-024-000-0280",
         "Phone 1": "555-0100", "Phone 1 DNC": "Y",
         "hcad_account": "0440240000280", "verdict": "PASS",
         "shape_flag": "SLIVER", "needs_manual_reason": None,
         "retail_estimate": None, "_asking": None},
        {"Address": "18 Oak Ln", "Zip": "77028", "APN": "044-024-000-0281",
         "Phone 1": "555-0101", "Phone 1 DNC": "N",
         "hcad_account": "0440240000281", "verdict": "NEGOTIATE",
         "retail_estimate": 80_000, "mao": 50_000, "needs_manual_reason": None,
         "_asking": 60_000.0},
        {"Address": "?? Unknown", "Zip": "77028", "APN": "",
         "hcad_account": None, "verdict": "",
         "needs_manual_reason": "missing APN", "_asking": None},
    ]


def test_write_xlsx_and_summary(tmp_path):
    rows = _sample_rows()
    headers = ["Address", "Zip", "APN", "Phone 1", "Phone 1 DNC"]
    xlsx = tmp_path / "enriched_leads.xlsx"
    write_xlsx(rows, headers, xlsx)

    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["leads", "needs_manual"]
    leads = wb["leads"]
    got_headers = [c.value for c in leads[1]]
    assert got_headers == headers + list(NEW_COLUMNS)
    assert leads.max_row == 3                        # header + 2 processed
    assert wb["needs_manual"].max_row == 2           # header + 1 reject
    # verdict fill: row 2 is the PASS row (red)
    pass_fill = leads.cell(row=2, column=1).fill.start_color.rgb
    assert pass_fill.endswith(VERDICT_FILLS["PASS"])
    # DNC column survived untouched
    dnc_idx = got_headers.index("Phone 1 DNC") + 1
    assert leads.cell(row=2, column=dnc_idx).value == "Y"

    summary = tmp_path / "summary.md"
    write_summary(rows, summary)
    text = summary.read_text()
    assert "**PASS**: 1" in text and "**NEGOTIATE**: 1" in text
    assert "needs_manual: 1" in text
    assert "18 Oak Ln" in text and "4210 Richland Dr" not in text.split("## Top")[1]


def test_render_svg(tmp_path):
    from .maps import render_svg

    row = {"Address": "4210 Richland Dr", "hcad_account": "0440240000280",
           "width_ft": 33.0, "depth_ft": 660.0, "aspect_ratio": 20.0,
           "verdict": "ASSEMBLAGE_LEAD"}
    roads = [{"name": "Richland Dr",
              "coords": from_local_feet([_road_ft], LON0, LAT0)[0]}]
    out = tmp_path / "p.svg"
    render_svg(SLIVER_RINGS, roads, row, out)
    svg = out.read_text()
    assert svg.startswith("<svg") and "Richland Dr" in svg
    assert VERDICT_FILLS["ASSEMBLAGE_LEAD"] in svg


# ---------------------------------------------------------------------------
# end-to-end pipeline (all seams stubbed)
# ---------------------------------------------------------------------------


def test_run_pipeline_end_to_end(tmp_path):
    from .cli import run_pipeline

    # marker order matters: the adjacency URL also mentions HCAD_NUM in
    # outFields, so match its unique geometryType first
    stub, calls = _arcgis_stub({
        "esriGeometryPolygon": HCAD_ADJ,
        "TxDOT_Roadways": TXDOT_ROADS,
        "NFHL": FEMA_X,
        "HCAD_NUM": HCAD_FEATURE,
    })

    def overpass_stub(method, url, payload, key):
        return 200, json.dumps(OVERPASS_ROAD)

    rows = [
        {"Address": "4210 Richland Dr", "Zip": "77028", "City": "Houston",
         "APN": "044-024-000-0280", "Lot Size Sqft": "21,780"},
        {"Address": "no apn here", "Zip": "77028", "APN": "",
         "Lot Size Sqft": "5000"},
    ]
    cache = Cache(tmp_path / "e2e.db")
    logs = []
    run_pipeline(rows, DEFAULT_CONFIG, http_request=stub,
                 overpass_request=overpass_stub, sleep=lambda s: None,
                 cache=cache, log=logs.append)

    sliver = rows[0]
    assert sliver["hcad_account"] == "0440240000280"
    assert sliver["shape_flag"] == "SLIVER"
    assert sliver["verdict"] == "ASSEMBLAGE_LEAD"     # LONE STAR HOMES LLC x2
    assert "LONE STAR HOMES LLC" in sliver["adjacent_owners"]
    assert rows[1]["needs_manual_reason"] == "missing APN"
    assert any("ASSEMBLAGE_LEAD" in line for line in logs)

    # resume: second run must not refetch (stage cache hits)
    before = len(calls)
    run_pipeline(rows, DEFAULT_CONFIG, http_request=stub,
                 overpass_request=overpass_stub, sleep=lambda s: None,
                 cache=cache, log=lambda m: None)
    assert len(calls) == before                       # zero new HTTP calls


def test_road_failures_retry_on_rerun_and_trip_breaker(tmp_path):
    from .cli import run_pipeline

    stub, _ = _arcgis_stub({
        "esriGeometryPolygon": HCAD_ADJ,
        "NFHL": FEMA_X,
        "HCAD_NUM": HCAD_FEATURE,
    })
    # disable the TxDOT primary so this exercises the Overpass fallback path
    cfg = DEFAULT_CONFIG.mutate(retry_delays_s=(0.1,), roads_arcgis_url="")
    cache = Cache(tmp_path / "roads.db")

    def fresh_rows():
        return [{"Address": f"{i} Richland Dr", "Zip": "77028",
                 "APN": "044-024-000-0280", "Lot Size Sqft": "21780"}
                for i in range(4)]

    down_calls = []

    def overpass_down(method, url, payload, key):
        down_calls.append(1)
        return 429, "overloaded"

    rows = fresh_rows()
    run_pipeline(rows, cfg, http_request=stub, overpass_request=overpass_down,
                 sleep=lambda s: None, cache=cache, log=lambda m: None)
    assert all(r["needs_manual_reason"] == "road data unavailable" for r in rows)
    # breaker: leads 1-3 each try 3 hosts x 2 attempts; lead 4 tries none
    assert len(down_calls) == 3 * 3 * 2

    # rerun with Overpass healthy: failures were NOT cached, so they retry
    rows2 = fresh_rows()
    run_pipeline(rows2, cfg, http_request=stub,
                 overpass_request=lambda *a: (200, json.dumps(OVERPASS_ROAD)),
                 sleep=lambda s: None, cache=cache, log=lambda m: None)
    assert all(not r["needs_manual_reason"] for r in rows2)
    assert all(r["verdict"] == "ASSEMBLAGE_LEAD" for r in rows2)  # sliver+builder


if __name__ == "__main__":
    import sys

    failures = 0
    module = sys.modules[__name__]
    for name in sorted(dir(module)):
        if name.startswith("test_"):
            fn = getattr(module, name)
            try:
                if "tmp_path" in fn.__code__.co_varnames[: fn.__code__.co_argcount]:
                    import tempfile
                    from pathlib import Path

                    with tempfile.TemporaryDirectory() as td:
                        fn(Path(td))
                else:
                    fn()
                print(f"ok {name}")
            except Exception as exc:  # noqa: BLE001
                failures += 1
                print(f"FAIL {name}: {exc}")
    raise SystemExit(1 if failures else 0)
